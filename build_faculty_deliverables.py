from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
PAPER_OUT = OUT / "paper"
FACULTY_OUT = OUT / "faculty"
PAPER_OUT.mkdir(parents=True, exist_ok=True)
FACULTY_OUT.mkdir(parents=True, exist_ok=True)
NAVY = "142C4E"
TEAL = "087E8B"
PALE = "EAF3F7"
GRAY = "536273"


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup_doc(doc, two_col=False):
    sec = doc.sections[0]
    sec.top_margin = Inches(.56 if two_col else .7)
    sec.bottom_margin = Inches(.52 if two_col else .7)
    sec.left_margin = Inches(.56 if two_col else .75)
    sec.right_margin = Inches(.56 if two_col else .75)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8.1 if two_col else 10.3)
    normal.paragraph_format.space_after = Pt(2 if two_col else 5)
    normal.paragraph_format.line_spacing = 0.99 if two_col else 1.08
    for nm, size, color in [("Heading 1", 10, NAVY), ("Heading 2", 9, TEAL)]:
        st = doc.styles[nm]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size if two_col else size + 1)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(4 if two_col else 9)
        st.paragraph_format.space_after = Pt(2)
    if two_col:
        sectpr = sec._sectPr
        cols = sectpr.xpath("./w:cols")
        if not cols:
            cols = [OxmlElement("w:cols")]
            sectpr.append(cols[0])
        cols[0].set(qn("w:num"), "2")
        cols[0].set(qn("w:space"), "360")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f = footer.add_run("AEROLAB  |  WORKING MANUSCRIPT — SIMULATION / SOFTWARE EVIDENCE ONLY" if two_col else "AEROLAB  |  FACULTY BRIEFING")
    f.font.name = "Arial"
    f.font.size = Pt(7)
    f.font.color.rgb = RGBColor.from_string(GRAY)


def p(doc, text="", bold_prefix=None, italic=False, align=None):
    par = doc.add_paragraph()
    if align is not None:
        par.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = par.add_run(bold_prefix)
        r.bold = True
        par.add_run(text[len(bold_prefix):])
    else:
        r = par.add_run(text)
        r.italic = italic
    return par


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def table(doc, headers, rows, widths=None, size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", size)
        shade(t.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, None, size)
            if ridx % 2 == 1:
                shade(cells[i], "F2F6F8")
    if widths:
        for row in t.rows:
            for idx, w in enumerate(widths):
                row.cells[idx].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def equation(doc, text):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before = Pt(1)
    q.paragraph_format.space_after = Pt(3)
    r = q.add_run(text)
    r.font.name = "Cambria Math"
    r.font.size = Pt(9)


def make_paper():
    doc = Document()
    setup_doc(doc, True)
    doc.sections[0].different_first_page_header_footer = True
    doc.core_properties.comments = "Working manuscript; claims limited to model/software evidence described in the text."
    # IEEE-style first page: title/abstract span the width, body uses two columns.
    cols = doc.sections[0]._sectPr.xpath("./w:cols")
    if cols:
        cols[0].set(qn("w:num"), "1")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    rr = title.add_run("AeroLab: An Embedded Control and Model-Based Fault-Monitoring Demonstrator")
    rr.bold = True
    rr.font.name = "Times New Roman"
    rr.font.size = Pt(17)
    rr.font.color.rgb = RGBColor.from_string(NAVY)
    doc.sections[0].first_page_footer.paragraphs[0].text = ""
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(1)
    r = author.add_run("Manav Sharma")
    r.bold = True
    r.font.size = Pt(10)
    aff = doc.add_paragraph("School of Electronics Engineering, Vellore Institute of Technology, Vellore, India")
    aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aff.paragraph_format.space_after = Pt(7)
    for run in aff.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    abstract = ("Abstract—This paper presents AeroLab, an ESP32-based control and telemetry demonstrator that combines a sampled-data linear-quadratic regulator (LQR), a Luenberger observer, a small multi-sensor interface, a browser dashboard, and MATLAB/Python visualization tools. The repository contains a second-order plant model parameterized by assumed natural frequency and damping ratio, MATLAB-generated discrete gains, firmware implementation, and software simulations of step response and injected model drift. The dashboard exposes telemetry and model-oriented traces, while firmware supports ultrasonic proximity sensing, a servo-mounted scanning sensor, inertial measurements, raw light-dependent-resistor (LDR) readings, and alarm controls. The principal contribution is an integrated, inspectable teaching platform and a reproducible software workflow—not a validated digital twin or a measured flight-control system. Code inspection shows that the fault residual compares internally propagated states and software-injected offsets rather than an independent physical actuator measurement; consequently, the current residual/CUSUM plots cannot establish hardware fault-detection performance. We document the architecture and equations, identify what the existing simulations support, and define a verification path based on calibrated sensors, synchronized independent measurements, logged experiments, and quantified detector performance. This evidence-aware framing makes the project suitable for a faculty demonstration while preventing simulation results from being misrepresented as hardware results.")
    ap = doc.add_paragraph()
    ap.paragraph_format.space_after = Pt(2)
    ar = ap.add_run(abstract)
    ar.font.size = Pt(8.1)
    ar.italic = False
    ip = doc.add_paragraph()
    ip.paragraph_format.space_after = Pt(6)
    ir = ip.add_run("Index Terms—embedded control, ESP32, linear-quadratic regulator, model-based monitoring, residual analysis, CUSUM, telemetry dashboard.")
    ir.italic = True
    ir.font.size = Pt(8.1)
    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    body_sec.top_margin = Inches(.56)
    body_sec.bottom_margin = Inches(.52)
    body_sec.left_margin = Inches(.56)
    body_sec.right_margin = Inches(.56)
    body_cols = body_sec._sectPr.xpath("./w:cols")
    if not body_cols:
        body_cols = [OxmlElement("w:cols")]
        body_sec._sectPr.append(body_cols[0])
    body_cols[0].set(qn("w:num"), "2")
    body_cols[0].set(qn("w:space"), "360")
    heading(doc, "I. INTRODUCTION")
    p(doc, "Digital-twin research spans a wide range of fidelity, from data-linked engineering models to continuously synchronized operational replicas. The term therefore needs a clear evidence boundary: a simulated model connected to an embedded interface is not, by itself, a validated digital twin [1], [2]. AeroLab is best understood at its present maturity as an embedded control and model-based monitoring demonstrator. It is useful for teaching how a control model, generated coefficients, firmware, sensors, and visualization can be assembled, while exposing the gap between software behavior and measured plant behavior.")
    p(doc, "This work describes the implementation visible in the project repository and distinguishes three evidence classes: (i) analytical model calculations, (ii) software-generated simulation traces, and (iii) physical telemetry. The project contains the first two and provides sensor interfaces for the third; the inspected sources do not provide an independently measured actuator-state data set suitable for validating the model residual. This distinction is central to reproducible engineering claims and to the presentation of fault-detection results [2], [8].")
    p(doc, "The contributions are: 1) a source-grounded description of the ESP32/MATLAB control workflow; 2) a precise interpretation of the dashboard scopes and simulation claims; and 3) a validation protocol that converts the demonstrator into a defensible experimental study. No flight-control or safety certification is claimed.")

    heading(doc, "II. RELATED WORK")
    p(doc, "The digital-twin paradigm has been discussed in aerospace engineering as a lifecycle link between physical assets and computational models [1]. Later surveys organize industrial digital-twin methods around model, data, and service integration [2]. Recent fault-diagnosis work likewise emphasizes that monitoring quality depends on meaningful measurements, residual construction, and validated decision rules [8]. These references motivate the architecture; they do not validate AeroLab's implementation.")
    p(doc, "State estimation using a model and measurement innovation follows the classical observer/filtering tradition [5]. Linear-quadratic regulation offers a systematic quadratic-cost feedback design, while discrete-time implementations must use gains and sample periods consistent with the sampled plant [6]. CUSUM originates in sequential change detection [3], with modern engineering guidance describing design and interpretation of control-chart thresholds [4]. Threshold selection trades detection speed against false alarms and must be evaluated under representative data; a threshold value alone is not a performance result.")

    heading(doc, "III. SYSTEM ARCHITECTURE")
    p(doc, "The project consists of four cooperating layers. First, control_simulation.m defines a second-order continuous model, discretizes it, computes LQR/observer gains, runs sampled-data simulations, and exports gnc_gains.h. Second, main.cpp executes on an ESP32 and manages sensors, a servo, alarms, state updates, and JSON telemetry. Third, dashboard.html and dashboard_analytics.html render the live/simulation interface and browser scopes. Fourth, Python and MATLAB scope applications provide alternate data visualization paths. The firmware loop uses a nominal sample period Ts = 0.05 s (20 Hz); actual packet cadence and deadline jitter require timestamped measurement rather than configuration alone.")
    p(doc, "The sensor panel includes fixed ultrasonic channel US1, servo-mounted ultrasonic channel US2, an MPU-6050, temperature and gas inputs, and an LDR analog input. Firmware reports the LDR as raw ADC data with ldrCalibrated=0. A single rotating range sensor measures range at successive azimuths; it cannot infer continuous two-dimensional hand pose. Thus the user-facing ‘hand tracking’ mode is more accurately described as a proximity-triggered scan/sector lock with reacquisition. Reliable operation depends on fresh echo validation, sensor alignment, electrical level shifting, power integrity, and common ground; the repository hardware checklist notes these dependencies.")

    heading(doc, "IV. MODEL AND CONTROL DESIGN")
    p(doc, "The MATLAB script parameterizes a second-order transfer function as follows:")
    equation(doc, "G(s) = ωₙ² / (s² + 2ζωₙs + ωₙ²),     ωₙ = 89 rad/s, ζ = 0.45.")
    p(doc, "These values are design assumptions in the checked-in script, not the output of a documented system-identification experiment. They imply the denominator coefficients 80.1 and 7921. In phase-variable form, x = [x₁, x₂]ᵀ with x₂ = ẋ₁. Zero-order-hold discretization at 50 ms yields x[k+1] = A_d x[k] + B_d u[k], y[k] = Cx[k]. The MATLAB script designs a discrete LQR using Q = diag(100, 1) and R = 0.01, and computes a discrete observer gain from selected continuous poles mapped to the sampled domain [5], [6].")
    equation(doc, "u[k] = sat₍₀,₁₎(N̄ r[k] − K x̂[k]),     x̂[k+1] = A_d x̂[k] + B_d u[k] + L(y[k] − Cx̂[k]).")
    p(doc, "The firmware header contains K = [−0.422821, 0.000667], N̄ = 0.577179, and L = [−0.186497, 8.481368]ᵀ. The discrete observer gain is designed for A_d and the mapped observer poles, but the firmware update in main.cpp applies the innovation correction to the state and then advances with a 1 ms propagation model between 50 ms controller updates. Thus the implemented observer is a multirate hybrid; it is not exactly the standard one-step 50 ms Luenberger recursion written above. The command is clipped to [0, 1]. An unsaturated linear LQR step response and the clipped firmware-like response are different experiments; neither is a physical servo measurement. Saturation, servo deadband, supply variation, load, and timing must be included before making actuator-performance claims.")

    heading(doc, "V. MODEL-BASED MONITORING AND DASHBOARD")
    p(doc, "The MATLAB software experiment propagates two internal states, injects an offset of 0.145 into the simulated measurement after 2.5 s, adds Gaussian noise with σ = 0.003, and evaluates a residual threshold δ = 0.100 and two-sided CUSUM with reference k = 0.02 and decision limit h = 0.06. The conceptual detector is:")
    equation(doc, "e[k] = yₘ[k] − yₜ[k];   S⁺[k] = max(0, S⁺[k−1] + e[k] − k);   S⁻[k] = max(0, S⁻[k−1] − e[k] − k).")
    p(doc, "The firmware source constructs y_measured from its internally propagated x1_real plus a software disturbance/drift term. Both x1_real and the twin state are propagated using the same model matrices and control input. The residual is therefore a software-model discrepancy under the injected scenario, not a difference between a calibrated physical actuator sensor and a model prediction. Although ultrasonic and inertial telemetry are physically read, they do not establish independent deflection feedback for this residual. Accordingly, all residual, CUSUM, Monte Carlo, and parameter-sweep figures in this version are simulation evidence only. CUSUM's literature basis [3], [4] does not convert synthetic test outcomes into measured false-alarm or detection guarantees.")
    p(doc, "The browser dashboard combines sensor indicators, alarms, control values, event logging, and scopes. To present results accurately, its data-source state must remain visible: DEMO/synthetic, software simulation, or live hardware. The twin coincidence trace is an internal model-state comparison; it is not independent plant/twin agreement. Raw LDR counts are not lux and cannot support quantitative illuminance claims until calibrated. Likewise, a range sensor can provide range-to-target at the current scan angle, not continuous hand tracking. These labels and limits are consistent with digital-twin practice, where synchronization and measurement provenance are essential [1], [2].")

    heading(doc, "VI. SOFTWARE EVALUATION")
    p(doc, "The checked-in MATLAB figures support analysis of the assumed model and deterministic/noisy simulations. Figure 1 compares continuous and discrete linear responses with a discrete response using firmware-like command saturation. Its purpose is to show the difference between ideal design analysis and a bounded command, not to report measured servo rise time. The Bode and root-locus figures characterize the specified model under their plotted feedback assumptions; they are not robustness certificates for the physical rig. The observer-error figure shows simulated convergence under the model, not estimation error against an independent sensor. The residual/CUSUM figure demonstrates response to an intentionally injected offset in software. The parameter sweep varies model parameters in simulation and is not an empirical reliability estimate.")
    p(doc, "The repository also includes repeated noisy runs and a nominal-controller/perturbed-model sweep. Such simulations are useful for debugging and sensitivity analysis, but the reported false-alarm fraction, detection fraction, or settling proportion—if generated—must be explicitly labeled as conditional on the selected synthetic noise, plant family, thresholds, initial state, and code version. They must not be reported as hardware rates or generalized outside that experiment. A paper-grade experimental result requires saved raw data, fixed test definitions, repeat counts, uncertainty intervals, and independent measurement.")

    heading(doc, "VII. REPRODUCIBILITY AND RECOMMENDED FIGURE SET")
    p(doc, "The model experiment is deterministic apart from seeded pseudo-random noise (rng(1) for the five-second injection run; separate seeds are assigned to the Monte Carlo and parameter sweep). Reproduction requires MATLAB Control System Toolbox-compatible functions including tf, ss, c2d, dlqr, place, stepinfo, margin, and exportgraphics. The generated firmware header is derived from the script and should be regenerated, then diffed, whenever model or design settings change. The Python scope test uses a deterministic synthetic scenario and establishes plotting/detector code behavior only; it does not validate hardware telemetry or the MATLAB numerical output.")
    p(doc, "For a concise conference paper, use three main figures: (1) the LQR benchmark with linear and saturated discrete traces clearly distinguished; (2) the synthetic injected-drift residual and normalized two-sided CUSUM, with injection and decision thresholds annotated; and (3) the architecture/data-provenance diagram distinguishing physical sensor channels from internally propagated model states. Move Bode, root-locus, observer simulation, and parameter sweep to supplementary material unless they answer a specific research question. Every figure caption should state ‘model simulation’ or ‘live hardware’ and report the test conditions. Do not caption the current model-state overlays as measured plant/twin agreement.")

    heading(doc, "VIII. LIMITATIONS AND VALIDATION PLAN")
    p(doc, "The main limitations are: (1) model parameters are assumed rather than identified; (2) the controlled plant state and twin are both propagated in software, so residual independence is absent; (3) no logged, synchronized actuator-position reference or calibrated sensor dataset is included; (4) LDR and gas sensor outputs require calibration and warm-up/conditioning; (5) single-sensor servo scanning supports sector-level proximity response, not general hand-pose tracking; and (6) the 20 Hz nominal loop does not by itself establish observed timing, packet freshness, or end-to-end latency. The ESP32 is an embedded demonstrator and not an aircraft-ready flight controller [7].")
    p(doc, "A defensible follow-on evaluation should first identify the physical servo/load response from repeated command-position trials, including deadband and saturation, and fit the model on one dataset while reserving separate trials for validation. Add an independent angular measurement (e.g., calibrated encoder or external optical reference), record synchronized command, measurement, IMU, ultrasonic validity, packet timestamp, and alarm state, and calculate sample-period jitter and data age. Calibrate LDR against a reference lux meter at multiple illumination levels. For hand interaction, define the target variable and workspace; a range-only rotating sensor should be evaluated as angular sector detection, while continuous tracking requires additional geometry/sensors or vision. Finally, tune residual/CUSUM thresholds on baseline training data and report false alarms per unit time, detection probability versus injected/measured fault magnitude, delay distribution, missed detections, and uncertainty across independent runs. Only then should the system be described as empirically validated monitoring.")

    heading(doc, "IX. CONCLUSION")
    p(doc, "AeroLab integrates model-based control design, ESP32 firmware, multi-sensor telemetry, browser visualization, and software fault-injection studies into a coherent educational demonstrator. Its strongest present takeaway is the end-to-end integration and its potential as a reproducible platform for measured validation. The current evidence supports model and simulation claims, not physical LQR tracking, hardware digital-twin fidelity, or field-ready fault-detection performance. Clearly distinguishing those evidence levels strengthens the technical presentation and provides a concrete path from prototype to a publishable experimental study.")

    heading(doc, "REFERENCES")
    refs = [
        "[1] E. H. Glaessgen and D. S. Stargel, “The digital twin paradigm for future NASA and U.S. Air Force vehicles,” in Proc. AIAA/ASME/ASCE/AHS/ASC Structures, Structural Dynamics, and Materials Conf., Honolulu, HI, USA, 2012, Art. no. AIAA 2012-1818, doi: 10.2514/6.2012-1818.",
        "[2] F. Tao, H. Zhang, A. Liu, and A. Y. C. Nee, “Digital twin in industry: State-of-the-art,” IEEE Trans. Ind. Informat., vol. 15, no. 4, pp. 2405–2415, Apr. 2019, doi: 10.1109/TII.2018.2873186.",
        "[3] E. S. Page, “Continuous inspection schemes,” Biometrika, vol. 41, nos. 1–2, pp. 100–115, Jun. 1954, doi: 10.1093/biomet/41.1-2.100.",
        "[4] NIST/SEMATECH, “CUSUM control charts,” e-Handbook of Statistical Methods, §6.3.2.3. [Online]. Available: https://www.itl.nist.gov/div898/handbook/pmc/section3/pmc323.htm. Accessed: Oct. 24, 2026.",
        "[5] R. E. Kalman, “A new approach to linear filtering and prediction problems,” J. Basic Eng., vol. 82, no. 1, pp. 35–45, Mar. 1960, doi: 10.1115/1.3662552.",
        "[6] MathWorks, “dlqr—Linear-quadratic regulator design for discrete-time systems,” Control System Toolbox Documentation. [Online]. Available: https://www.mathworks.com/help/control/ref/lti.dlqr.html. Accessed: Oct. 24, 2026.",
        "[7] Espressif Systems, “ESP32 series datasheet,” Version 5.2. [Online]. Available: https://documentation.espressif.com/esp32_datasheet_en.html. Accessed: Oct. 24, 2026.",
        "[8] A. C. S. Araujo, L. M. Rabelo, P. C. M. Lamim Filho, and W. M. Caminhas, “A digital twin approach to smart monitoring and fault diagnosis,” IEEE Access, vol. 13, pp. 148384–148395, 2025, doi: 10.1109/ACCESS.2025.3601349."
    ]
    for ref in refs:
        par = doc.add_paragraph(ref)
        par.paragraph_format.left_indent = Inches(.12)
        par.paragraph_format.first_line_indent = Inches(-.12)
        par.paragraph_format.space_after = Pt(2)
        for run in par.runs:
            run.font.size = Pt(6.8)
    doc.core_properties.title = "AeroLab: An Embedded Control and Model-Based Fault-Monitoring Demonstrator"
    doc.core_properties.subject = "IEEE-style working manuscript with explicit evidence limitations"
    doc.core_properties.author = "Manav Sharma"
    doc.save(PAPER_OUT / "AeroLab_IEEE_Style_Manuscript.docx")


def make_guide():
    doc = Document()
    setup_doc(doc)
    top = doc.add_paragraph()
    top.paragraph_format.space_after = Pt(2)
    rr = top.add_run("AEROLAB")
    rr.bold = True
    rr.font.size = Pt(11)
    rr.font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run("Faculty Demonstration & Graph Walkthrough")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    heading(doc, "The one-sentence takeaway")
    call = doc.add_table(rows=1, cols=1)
    call.style = "Table Grid"
    cell = call.cell(0, 0)
    shade(cell, PALE)
    set_cell_text(cell, "AeroLab connects ESP32 control, sensors, a dashboard, and software fault simulations. Today’s graphs support model/simulation claims; independent hardware validation is the next research step.", True, NAVY, 10.5)
    heading(doc, "Suggested 5-minute faculty walkthrough")
    table(doc, ["Time", "Show / explain", "Suggested wording"], [
        ["0:00–0:40", "Dashboard source: DEMO, simulation, or live.", "I’ll distinguish simulation from physical telemetry."],
        ["0:40–1:25", "MATLAB → ESP32 → sensors/actuator → dashboard.", "The contribution is an integrated, auditable teaching demonstrator."],
        ["1:25–2:10", "US1 fixed, US2 servo-mounted; echo validity and raw LDR.", "This is angular-sector proximity response, not continuous hand pose."],
        ["2:10–3:15", "Step plot: linear models versus clipped response.", "This is a model benchmark, not measured servo performance."],
        ["3:15–4:05", "Observer/residual/CUSUM with injected-fault marker.", "The offset is injected in software; residual is not independent sensor validation."],
        ["4:05–5:00", "State the next validation experiment.", "I’ll measure servo angle independently and report repeatable timing/detector statistics."],
    ], [0.8, 2.55, 3.15], 8.2)
    heading(doc, "Terminology guardrails")
    table(doc, ["Say", "Avoid until measured"], [
        ["Model-linked demonstrator", "Validated digital twin / flight controller"],
        ["Servo ultrasonic sector scan", "Continuous 2-D hand tracking"],
        ["Software-injected drift", "Hardware fault-detection validation"],
        ["Raw, uncalibrated LDR ADC", "Lux / calibrated photometry"],
        ["Nominal 20 Hz loop", "Verified end-to-end timing"],
        ["Simulated parameter sweep", "Guaranteed hardware robustness"],
    ], [3.25, 3.25], 8.0)
    heading(doc, "How to read each graph")

    table(doc, ["Graph / scope", "What to point out", "What it means / caveat"], [
        ["Step-response comparison", "Setpoint and continuous, discrete-linear, and saturated curves.", "Compare rise, overshoot, settling, and input bounds. Saturated trace is closer to firmware command logic, but remains model-only; do not quote it as measured servo dynamics."],
        ["Twin/model-state coincidence", "Overlapping x₁ and twin-state traces; divergence when software offset is introduced.", "Internal states share model propagation/control. Agreement is not independent validation of a physical plant."],
        ["Residual and CUSUM", "Residual magnitude against δ; positive/negative accumulators against h; injection time.", "Illustrates detector logic under a chosen synthetic fault/noise. It does not establish hardware false-alarm probability or detection latency."],
        ["Control effort and servo", "Normalized u in [0,1] and servo-angle command.", "Shows commanded effort/angle mapping. Distinguish command from independently measured shaft position."],
        ["Observer estimation error", "Position error and scaled velocity error decay in simulated run.", "Model observer behavior only; no physical reference measurement is present."],
        ["Bode / root locus", "Frequency response and nominal model pole movement under plotted assumptions.", "Useful design analysis for the assumed model; not an empirical stability margin or robustness certificate for the hardware."],
        ["Parameter sweep", "Spread of model responses under the stated parameter perturbations.", "Sensitivity study on simulated plants. Do not state as measured reliability or a guarantee."],
        ["Attitude / sensor scopes", "Roll/pitch, vibration, distance, temperature/gas, raw LDR.", "Live only when source and freshness indicate hardware. LDR is ADC counts until calibrated; ultrasonic channels require fresh valid echoes."],
    ], [1.45, 2.3, 2.75], 8)

    heading(doc, "Dashboard cleanup before the meeting")
    for line in [
        "Keep the source indicator visible; never present DEMO/simulation as LIVE.",
        "Use corrected analytics: assumed model parameters, firmware-matched gains, and no unsupported FAR/robustness claims.",
        "Call MATLAB views model analysis. Label residual/CUSUM as software-model injected simulation.",
        "For hardware demo, verify fresh packet age and sensor validity; record CSV. Connection alone is not validation.",
        "Describe hand mode as sector scan/lock; verify US2 echo, servo motion, reacquisition, alarms, and limits.",
        "Call LDR/MQ raw or uncalibrated unless calibration and warm-up are documented.",
    ]:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(2)
        par.add_run(line)

    heading(doc, "Likely faculty questions")
    table(doc, ["Question", "Short answer"], [
        ["Is it a digital twin?", "Not yet validated: residual uses internal model states plus injected drift, not an independent actuator measurement."],
        ["Does it track a hand?", "It scans and locks/reacquires a range sector; one sensor cannot recover continuous 2-D hand pose."],
        ["Where did ωₙ, ζ come from?", "They are design assumptions; identify them from measured servo/load data."],
        ["What does the alarm prove?", "Software response to coded synthetic offset only; hardware rates remain unknown."],
        ["Strongest result?", "Integrated model-to-firmware-to-dashboard teaching demonstrator."],
        ["Next experiment?", "Measure shaft angle independently, synchronize logs, identify/validate model, test held-out faults."],
    ], [1.65, 4.85], 8.0)

    validation_heading = heading(doc, "A strong next validation package")
    table(doc, ["Validation area", "Measure before a performance claim"], [
        ["Hardware provenance", "Board and firmware version, supply voltage, wiring, sensor identity, environment, calibration procedure, and test date."],
        ["Independent actuator ground truth", "Encoder or optical angle reference synchronized with command and telemetry timestamps."],
        ["Model identification", "Repeated step/frequency tests; fit residuals and validation on separate runs, including deadband, rate limits, and saturation."],
        ["Runtime timing", "Packet age, missed packets, loop-period/jitter distribution, command-to-measurement delay, and reset/error logs."],
        ["Fault detector", "Baseline/fault scenarios; false alarms per hour, detection probability vs fault magnitude, delay distribution, misses, sample count, and confidence intervals."],
        ["Photonic sensor", "LDR ADC-to-lux calibration against a reference meter at multiple levels; document geometry and ambient-light control."],
        ["Hand interaction", "Define target distance/angle/speed and score sector detection/reacquisition. Continuous 2-D tracking needs adequate additional sensing geometry."],
    ], [1.45, 5.05], 8.0)
    doc.core_properties.title = "AeroLab Faculty Demonstration and Graph Walkthrough"
    doc.core_properties.subject = "Project presentation guide and evidence boundaries"
    doc.core_properties.author = "Manav Sharma"
    doc.save(FACULTY_OUT / "AeroLab_Faculty_Demo_Guide_current.docx")


if __name__ == "__main__":
    make_paper()
    make_guide()
    print("Created:")
    print(PAPER_OUT / "AeroLab_IEEE_Style_Manuscript.docx")
    print(FACULTY_OUT / "AeroLab_Faculty_Demo_Guide_current.docx")
